import hashlib
import io
import threading
import time
import unittest
import zipfile
from pathlib import Path
from unittest.mock import MagicMock, patch

from docx import Document
from pypdf import PdfWriter

from deepdoc_http_service import DeepDocRuntime, create_app


def pdf_bytes() -> bytes:
    output = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(output)
    return output.getvalue()


def encrypted_pdf_bytes(password: str) -> bytes:
    output = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.encrypt(password)
    writer.write(output)
    return output.getvalue()


def docx_bytes() -> bytes:
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", "<document/>")
    return output.getvalue()


class FakeParser:
    def __init__(self) -> None:
        self.active = 0
        self.max_active = 0
        self.lock = threading.Lock()
        self.outlines = []
        self.boxes = [
            {
                "text": "Document title",
                "layout_type": "title",
                "positions": [[1, 1, 2, 3, 4]],
            },
            {
                "text": "A😀B",
                "layout_type": "text",
                "positions": [[1, 1, 2, 5, 6]],
            },
            {
                "text": "a | b",
                "layout_type": "table",
                "positions": [[1, 1, 2, 7, 8]],
            },
        ]

    def parse_into_bboxes(self, *_args, **_kwargs):
        with self.lock:
            self.active += 1
            self.max_active = max(self.max_active, self.active)
        try:
            time.sleep(0.08)
            return self.boxes
        finally:
            with self.lock:
                self.active -= 1


class FakeDocxParser:
    def __call__(self, _path):
        return (
            [
                ("Document title", "Title"),
                ("Section", "Heading 2"),
                ("Body", "Normal"),
            ],
            [["Name: A;Value: 1"]],
        )


class FakeTxtParser:
    def __call__(self, _path, **_kwargs):
        return [["first\nsecond", ""], ["third", ""]]


class FakeMarkdownExtractor:
    def __init__(self, _text):
        pass

    def extract_elements(self, include_meta=False):
        assert include_meta
        return [
            {
                "type": "header",
                "content": "# Heading",
                "start_line": 0,
                "end_line": 0,
            },
            {
                "type": "text_block",
                "content": "Body",
                "start_line": 1,
                "end_line": 1,
            },
        ]


class DeepDocHttpServiceTest(unittest.TestCase):
    def setUp(self) -> None:
        self.parser = FakeParser()
        self.runtime = DeepDocRuntime(
            {
                "PDF": self.parser,
                "DOCX": FakeDocxParser(),
                "TXT": FakeTxtParser(),
                "MARKDOWN": FakeMarkdownExtractor,
            },
            {"parser.fake": ["CUDAExecutionProvider"]},
            max_pdf_pages=10,
            max_blocks=100,
            parse_concurrency=1,
        )
        self.app = create_app(self.runtime, 1024 * 1024)
        self.payload = pdf_bytes()
        self.digest = hashlib.sha256(self.payload).hexdigest()

    def post(self, payload=None, source_format="PDF"):
        payload = self.payload if payload is None else payload
        digest = hashlib.sha256(payload).hexdigest()
        with self.app.test_client() as client:
            return client.post(
                "/v1/parse",
                data={"file": (io.BytesIO(payload), "document.bin")},
                headers={
                    "X-Source-SHA256": digest,
                    "X-Source-Format": source_format,
                },
                content_type="multipart/form-data",
            )

    def test_maps_neutral_blocks_and_utf16_offsets(self):
        response = self.post()
        self.assertEqual(200, response.status_code)
        body = response.get_json()
        self.assertEqual("docquery-deepdoc-http-v1", body["schemaVersion"])
        self.assertEqual("PDF", body["sourceFormat"])
        self.assertEqual("TITLE", body["blocks"][0]["kind"])
        self.assertTrue(body["blocks"][0]["documentTitle"])
        self.assertEqual("PARAGRAPH", body["blocks"][1]["kind"])
        self.assertEqual(4, body["blocks"][1]["pageCharacterEnd"]
                         - body["blocks"][1]["pageCharacterStart"])
        self.assertEqual("TABLE_CELL", body["blocks"][2]["kind"])
        self.assertIn(
            "DEEPDOC_TABLE_STRUCTURE_FLATTENED",
            {warning["code"] for warning in body["warnings"]},
        )

    def test_expands_html_table_into_cells_with_grid_coordinates(self):
        self.parser.boxes[-1]["text"] = (
            "<table><tr><th rowspan='2'>Name</th><th>Value</th></tr>"
            "<tr><td>A &amp; B</td></tr></table>"
        )

        response = self.post()

        self.assertEqual(200, response.status_code)
        body = response.get_json()
        cells = [block for block in body["blocks"] if block["kind"] == "TABLE_CELL"]
        self.assertEqual(["Name", "Value", "A & B"], [cell["text"] for cell in cells])
        self.assertEqual(
            [(0, 0), (0, 1), (1, 1)],
            [(cell["tableRow"], cell["tableColumn"]) for cell in cells],
        )
        self.assertEqual({"pdf:p1:t0"}, {cell["tableId"] for cell in cells})
        self.assertEqual([2, 1, 1], [cell["tableRowSpan"] for cell in cells])
        self.assertEqual(
            ["Name", "Value", "Value"],
            [cell["tableColumnHeader"] for cell in cells],
        )
        warning_codes = {warning["code"] for warning in body["warnings"]}
        self.assertIn("DEEPDOC_TABLE_CELLS_STRUCTURED", warning_codes)
        self.assertIn("DEEPDOC_TABLE_COLUMN_HEADERS_BOUND", warning_codes)
        self.assertNotIn("DEEPDOC_TABLE_STRUCTURE_FLATTENED", warning_codes)

    def test_binds_aligned_visual_titles_to_table_columns_without_changing_text(self):
        boxes = [
            {
                "text": "<table><tr><td>Earlier continuation</td><td></td></tr>"
                        "<tr><td>Quarter three body</td>"
                        "<td>Quarter four body</td></tr></table>",
                "layout_type": "table",
                "positions": [[16, 30, 570, 10, 715]],
            },
            {
                "text": "QUARTER 3:",
                "layout_type": "title",
                "positions": [[16, 40, 110, 146, 158]],
            },
            {
                "text": "QUARTER 4:",
                "layout_type": "title",
                "positions": [[16, 310, 380, 159, 171]],
            },
            {
                "text": "Later bold label",
                "layout_type": "title",
                "positions": [[16, 40, 180, 420, 432]],
            },
        ]

        blocks, warnings = self.runtime._map_pdf_boxes(boxes, 16, [])

        cells = [block for block in blocks if block["kind"] == "TABLE_CELL"]
        self.assertEqual(
            ["Earlier continuation", "Quarter three body", "Quarter four body"],
            [cell["text"] for cell in cells],
        )
        self.assertEqual(
            [None, "QUARTER 3:", "QUARTER 4:"],
            [cell["tableColumnHeader"] for cell in cells],
        )
        self.assertEqual({"pdf:p16:t0"}, {cell["tableId"] for cell in cells})
        self.assertIn(
            "DEEPDOC_TABLE_COLUMN_HEADERS_BOUND",
            {warning["code"] for warning in warnings},
        )

    def test_prefers_ordered_digital_table_parent_without_duplicate_title_blocks(self):
        boxes = [
            {
                "text": "<table><tr><td>Google body Small business body</td>"
                        "<td>Quarter four body</td></tr></table>",
                "layout_type": "table",
                "positions": [[16, 30, 570, 10, 715]],
            },
            {
                "text": "QUARTER 3:",
                "layout_type": "title",
                "positions": [[16, 40, 110, 146, 158]],
            },
            {
                "text": "QUARTER 4:",
                "layout_type": "title",
                "positions": [[16, 310, 380, 159, 171]],
            },
            {
                "text": "Preparing for Tomorrow's Workplace Skills",
                "layout_type": "title",
                "positions": [[16, 40, 247, 420, 432]],
            },
        ]
        ordered_cells = self.runtime._pdf_table_matrix_cells(
            [[
                "QUARTER 3:\nPreparing for Tomorrow's Workplace Skills\n"
                "Use Google\nPreparing for Tomorrow's Workplace Skills\n"
                "Pick a small business",
                "QUARTER 4:\nQuarter four body",
            ]]
        )

        with patch.object(
            self.runtime,
            "_pdf_digital_table_cells",
            return_value=({0: ordered_cells}, {1, 2, 3}),
        ):
            blocks, warnings = self.runtime._map_pdf_boxes(
                boxes,
                16,
                [],
                Path("unused.pdf"),
            )

        self.assertEqual(2, len(blocks))
        self.assertTrue(all(block["kind"] == "TABLE_CELL" for block in blocks))
        self.assertEqual(
            "QUARTER 3:\nPreparing for Tomorrow's Workplace Skills\n"
            "Use Google\nPreparing for Tomorrow's Workplace Skills\n"
            "Pick a small business",
            blocks[0]["text"],
        )
        self.assertEqual("QUARTER 3:", blocks[0]["tableColumnHeader"])
        self.assertEqual("QUARTER 4:", blocks[1]["tableColumnHeader"])
        self.assertIn(
            "DEEPDOC_DIGITAL_TABLE_ORDER_PRESERVED",
            {warning["code"] for warning in warnings},
        )

    def test_detects_digital_tables_once_per_page_including_empty_results(self):
        boxes = [
            {"text": "First table", "positions": [[1, 0, 100, 0, 10]]},
            {"text": "Second table", "positions": [[1, 0, 100, 20, 30]]},
            {"text": "Third table", "positions": [[2, 0, 100, 0, 10]]},
        ]
        structured_tables = {
            index: self.runtime._pdf_table_matrix_cells([[box["text"]]])
            for index, box in enumerate(boxes)
        }
        for detected in (True, False):
            with self.subTest(detected=detected):
                candidates = []
                for box in boxes:
                    _, left, right, top, bottom = box["positions"][0]
                    candidate = MagicMock(bbox=(left, top, right, bottom))
                    candidate.extract.return_value = [[box["text"]]]
                    candidates.append(candidate)
                pages = [MagicMock(), MagicMock()]
                pages[0].find_tables.return_value = candidates[:2] if detected else []
                pages[1].find_tables.return_value = candidates[2:] if detected else []
                pdfplumber = MagicMock()
                pdfplumber.open.return_value.__enter__.return_value.pages = pages

                with (
                    patch.dict("sys.modules", {"pdfplumber": pdfplumber}),
                    patch.object(
                        self.runtime,
                        "_pdf_table_candidate_cells",
                        side_effect=lambda _page, _candidate, matrix:
                            self.runtime._pdf_table_matrix_cells(matrix),
                    ),
                    patch.object(
                        self.runtime, "_pdf_relational_table_cells", return_value=[]
                    ),
                ):
                    replacements, covered = self.runtime._pdf_digital_table_cells(
                        Path("unused.pdf"), boxes, structured_tables
                    )

                for page in pages:
                    page.find_tables.assert_called_once_with()
                self.assertEqual(structured_tables if detected else {}, replacements)
                self.assertEqual(set(), covered)
                if detected:
                    for candidate in candidates:
                        candidate.extract.assert_called_once_with()

    def test_splits_ordered_table_column_into_title_body_children(self):
        cells = self.runtime._pdf_table_line_cells(
            [
                ("REGION A", True, False),
                ("First activity", True, False),
                ("Interview", True, False),
                ("Interview classmates.", False, False),
                ("Second activity", True, False),
                ("Use the reference site.", False, False),
                ("Third activity", True, False),
                ("Pick a small business.", False, False),
                ("Fourth activity", True, False),
                ("Pick a specific product.", False, False),
                ("Resource task - Review the directory.", True, True),
                ("Quiz - Comprehension", True, True),
            ],
            0,
            0,
        )

        self.assertEqual(
            [
                "REGION A\nFirst activity\nInterview\nInterview classmates.",
                "Second activity\nUse the reference site.",
                "Third activity\nPick a small business.",
                "Fourth activity\nPick a specific product.",
                "Resource task - Review the directory.",
                "Quiz - Comprehension",
            ],
            [cell["text"] for cell in cells],
        )
        self.assertEqual(
            [
                "REGION A\nFirst activity\nInterview",
                "Second activity",
                "Third activity",
                "Fourth activity",
                "Resource task - Review the directory.",
                "Quiz - Comprehension",
            ],
            [cell["headingText"] for cell in cells],
        )
        self.assertTrue(all(cell["column"] == 0 for cell in cells))

    def test_reconstructs_multi_row_two_column_table_from_word_geometry(self):
        class FakePage:
            def __init__(self, words):
                self.words = words

            def crop(self, _bbox):
                return self

            def extract_words(self, **_kwargs):
                return self.words

        words = []
        for top, left_text, right_text in [
            (10, "Total", "59%"),
            (20, "Gender", None),
            (30, "Male", "65"),
            (40, "Female", "55"),
        ]:
            words.append({"text": left_text, "top": top, "x0": 10, "x1": 55})
            if right_text:
                words.append(
                    {"text": right_text, "top": top, "x0": 150, "x1": 175}
                )
        deepdoc_cells = [
            {"row": 0, "column": 0, "text": "Total Gender Male"},
            {"row": 0, "column": 1, "text": "59% 65"},
            {"row": 1, "column": 0, "text": "Female"},
            {"row": 1, "column": 1, "text": "55"},
            {"row": 2, "column": 0, "text": "end"},
        ]

        cells = self.runtime._pdf_relational_table_cells(
            FakePage(words),
            (0, 0, 200, 60),
            deepdoc_cells,
        )

        self.assertEqual(
            [
                (0, 0, "Total"),
                (0, 1, "59%"),
                (1, 0, "Gender"),
                (2, 0, "Male"),
                (2, 1, "65"),
                (3, 0, "Female"),
                (3, 1, "55"),
            ],
            [(cell["row"], cell["column"], cell["text"]) for cell in cells],
        )

    def test_does_not_reconstruct_one_row_column_flow_as_relational_table(self):
        class FailIfReadPage:
            def crop(self, _bbox):
                raise AssertionError("one-row column flow must not inspect word geometry")

        cells = self.runtime._pdf_relational_table_cells(
            FailIfReadPage(),
            (0, 0, 200, 100),
            [
                {"row": 0, "column": 0, "text": "Left narrative"},
                {"row": 0, "column": 1, "text": "Right narrative"},
            ],
        )

        self.assertEqual([], cells)

    def test_binds_caption_row_label_and_item_heading_as_table_metadata(self):
        boxes = [
            {
                "text": "Adoption by gender",
                "layout_type": "text",
                "positions": [[1, 10, 90, 10, 18]],
            },
            {
                "text": (
                    "<table><tr><th>Group</th><th>Rate</th></tr>"
                    "<tr><td>Male</td><td>65%</td></tr></table>"
                ),
                "layout_type": "table",
                "positions": [[1, 10, 90, 24, 70]],
            },
        ]

        blocks, warnings = self.runtime._map_pdf_boxes(boxes, 1, [])

        cells = [block for block in blocks if block["kind"] == "TABLE_CELL"]
        male = next(cell for cell in cells if cell["text"] == "Male")
        rate = next(cell for cell in cells if cell["text"] == "65%")
        self.assertEqual("Adoption by gender", rate["tableCaption"])
        self.assertEqual("Male", male["tableRowHeader"])
        self.assertEqual("Male", rate["tableRowHeader"])
        self.assertEqual("Rate", rate["tableColumnHeader"])
        self.assertEqual("pdf:p1:t0", rate["tableGroupId"])
        self.assertIsNone(rate["tableContinuationOf"])
        self.assertEqual("PARAGRAPH", blocks[0]["kind"])
        self.assertIn(
            "DEEPDOC_TABLE_RELATIONS_BOUND",
            {warning["code"] for warning in warnings},
        )

    def test_ignores_header_cells_outside_the_first_table_row(self):
        cells = [
            {"row": 0, "column": 0, "text": "Male", "header": False},
            {"row": 0, "column": 1, "text": "65", "header": False},
            {"row": 1, "column": 0, "text": "Female", "header": True},
            {"row": 1, "column": 1, "text": "55", "header": True},
        ]

        self.assertEqual({}, self.runtime._html_table_column_headers(cells, 2))

    def test_does_not_guess_unaligned_visual_table_headers(self):
        boxes = [
            {
                "text": "<table><tr><td>Left</td><td>Right</td></tr></table>",
                "layout_type": "table",
                "positions": [[1, 0, 100, 0, 100]],
            },
            {
                "text": "Possible left label",
                "layout_type": "title",
                "positions": [[1, 5, 40, 10, 18]],
            },
            {
                "text": "Possible right label",
                "layout_type": "title",
                "positions": [[1, 55, 95, 80, 88]],
            },
        ]

        blocks, warnings = self.runtime._map_pdf_boxes(boxes, 1, [])

        cells = [block for block in blocks if block["kind"] == "TABLE_CELL"]
        self.assertTrue(all(cell["tableColumnHeader"] is None for cell in cells))
        self.assertEqual({"pdf:p1:t0"}, {cell["tableId"] for cell in cells})
        self.assertNotIn(
            "DEEPDOC_TABLE_COLUMN_HEADERS_BOUND",
            {warning["code"] for warning in warnings},
        )

    def test_two_requests_wait_for_one_parse_slot_without_429(self):
        statuses = []

        def call():
            statuses.append(self.post().status_code)

        first = threading.Thread(target=call)
        second = threading.Thread(target=call)
        first.start()
        second.start()
        first.join()
        second.join()

        self.assertEqual([200, 200], sorted(statuses))
        self.assertEqual(1, self.parser.max_active)
        self.assertEqual(1, self.runtime.max_observed_active)

    def test_accepts_pdf_with_permission_encryption_and_empty_password(self):
        response = self.post(encrypted_pdf_bytes(""))

        self.assertEqual(200, response.status_code)

    def test_rejects_pdf_that_requires_a_password(self):
        response = self.post(encrypted_pdf_bytes("secret"))

        self.assertEqual(422, response.status_code)
        self.assertEqual(
            "ENCRYPTED_DOCUMENT_UNSUPPORTED",
            response.get_json()["error"]["code"],
        )

    def test_governs_title_candidates_without_dropping_their_text(self):
        boxes = [
            {"text": "Document title", "layout_type": "title", "positions": [[1, 0, 90, 0, 8]]},
            {"text": "UNIT 1: Foundations", "layout_type": "title", "positions": [[1, 0, 90, 15, 23]]},
            {"text": "Substantial body text that belongs to the numbered unit.", "layout_type": "text", "positions": [[1, 0, 90, 25, 35]]},
            {"text": "Repeated label", "layout_type": "title", "positions": [[1, 0, 90, 40, 48]]},
            {"text": "First repeated-label body remains searchable.", "layout_type": "text", "positions": [[1, 0, 90, 50, 60]]},
            {"text": "Repeated label", "layout_type": "title", "positions": [[2, 0, 90, 40, 48]]},
            {"text": "Second repeated-label body remains searchable.", "layout_type": "text", "positions": [[2, 0, 90, 50, 60]]},
            {"text": "Empty visual title", "layout_type": "title", "positions": [[2, 0, 90, 70, 78]]},
            {"text": "APPENDIX A: Evidence", "layout_type": "title", "positions": [[2, 0, 90, 80, 88]]},
        ]

        blocks, warnings = self.runtime._map_pdf_boxes(
            boxes,
            2,
            [("UNIT 1: Foundations", 1, 1)],
        )

        by_text = {block["text"]: block for block in blocks}
        self.assertEqual("TITLE", by_text["Document title"]["kind"])
        self.assertEqual("HEADING", by_text["UNIT 1: Foundations"]["kind"])
        self.assertEqual(2, by_text["UNIT 1: Foundations"]["headingLevel"])
        self.assertEqual("PDF_OUTLINE_MATCH", by_text["UNIT 1: Foundations"]["detectionSource"])
        self.assertEqual("PARAGRAPH", by_text["Repeated label"]["kind"])
        self.assertIsNone(by_text["Repeated label"]["headingLevel"])
        self.assertEqual("PARAGRAPH", by_text["Empty visual title"]["kind"])
        self.assertEqual("HEADING", by_text["APPENDIX A: Evidence"]["kind"])
        self.assertEqual(len(boxes), len(blocks))
        codes = {warning["code"] for warning in warnings}
        self.assertIn("DEEPDOC_TITLE_CANDIDATES_GOVERNED", codes)
        self.assertIn("DEEPDOC_REPEATED_TITLE_DOWNGRADED", codes)
        self.assertIn("DEEPDOC_EMPTY_TITLE_DOWNGRADED", codes)
        self.assertIn("DEEPDOC_TITLE_GOVERNANCE_TEXT_PRESERVED", codes)

    def test_merges_only_adjacent_weak_title_lines_on_the_same_page(self):
        boxes = [
            {"text": "Document title", "layout_type": "title", "positions": [[1, 0, 90, 0, 8]]},
            {"text": "Board meetings during the financial", "layout_type": "title", "positions": [[1, 0, 80, 20, 28]]},
            {"text": "year", "layout_type": "title", "positions": [[1, 0, 25, 30, 38]]},
            {"text": "Body text remains after the merged heading and is long enough.", "layout_type": "text", "positions": [[1, 0, 90, 45, 55]]},
            {"text": "UNIT 2: Strong heading", "layout_type": "title", "positions": [[1, 0, 90, 60, 68]]},
            {"text": "I.", "layout_type": "title", "positions": [[1, 0, 20, 70, 78]]},
        ]

        blocks, warnings = self.runtime._map_pdf_boxes(boxes, 1, [])

        texts = [block["text"] for block in blocks]
        self.assertIn("Board meetings during the financial year", texts)
        self.assertNotIn("Board meetings during the financial", texts)
        self.assertNotIn("year", texts)
        self.assertIn("UNIT 2: Strong heading", texts)
        self.assertIn("I.", texts)
        self.assertEqual("HEADING", next(block for block in blocks if block["text"] == "I.")["kind"])
        self.assertEqual(5, len(blocks))
        self.assertIn(
            "DEEPDOC_TITLE_FRAGMENTS_MERGED",
            {warning["code"] for warning in warnings},
        )

    def test_merges_inline_number_and_title_without_merging_numbered_body(self):
        boxes = [
            {"text": "Document title", "layout_type": "title", "positions": [[1, 0, 90, 0, 8]]},
            {"text": "4.1.", "layout_type": "title", "positions": [[2, 10, 22, 20, 28]]},
            {"text": "Payment Elections.", "layout_type": "title", "positions": [[2, 28, 85, 20, 28]]},
            {"text": "Section body remains searchable and provides enough evidence.", "layout_type": "text", "positions": [[2, 10, 90, 34, 44]]},
            {"text": "1.", "layout_type": "title", "positions": [[2, 10, 22, 55, 63]]},
            {"text": "A numbered body item must not become a heading title.", "layout_type": "text", "positions": [[2, 28, 90, 55, 63]]},
        ]

        blocks, warnings = self.runtime._map_pdf_boxes(boxes, 2, [])

        by_text = {block["text"]: block for block in blocks}
        self.assertEqual("HEADING", by_text["4.1. Payment Elections."]["kind"])
        self.assertEqual(2, by_text["4.1. Payment Elections."]["headingLevel"])
        self.assertNotIn("4.1.", by_text)
        self.assertEqual("HEADING", by_text["1."]["kind"])
        self.assertIn("A numbered body item must not become a heading title.", by_text)
        self.assertIn(
            "DEEPDOC_TITLE_FRAGMENTS_MERGED",
            {warning["code"] for warning in warnings},
        )

    def test_does_not_merge_numbered_title_with_title_on_another_line(self):
        boxes = [
            {"text": "Document title", "layout_type": "title", "positions": [[1, 0, 90, 0, 8]]},
            {"text": "4.1.", "layout_type": "title", "positions": [[2, 10, 22, 20, 28]]},
            {"text": "Payment Elections.", "layout_type": "title", "positions": [[2, 10, 85, 34, 42]]},
            {"text": "Body text remains available after both candidates.", "layout_type": "text", "positions": [[2, 10, 90, 48, 58]]},
        ]

        blocks, _warnings = self.runtime._map_pdf_boxes(boxes, 2, [])

        texts = [block["text"] for block in blocks]
        self.assertIn("4.1.", texts)
        self.assertIn("Payment Elections.", texts)
        self.assertNotIn("4.1. Payment Elections.", texts)

    def test_reconstructs_numbered_heading_from_same_row_text_fragment(self):
        boxes = [
            {"text": "Document title", "layout_type": "title", "positions": [[1, 0, 90, 0, 8]]},
            {"text": "4.1.", "layout_type": "title", "positions": [[2, 10, 22, 20, 28]]},
            {"text": "Payment Elections.", "layout_type": "text", "positions": [[2, 28, 85, 20, 28]]},
            {"text": "4.3.", "layout_type": "text", "positions": [[2, 10, 22, 40, 48]]},
            {"text": "Deferred Cash. For each director, the company records deferred cash.", "layout_type": "text", "positions": [[2, 28, 95, 40, 48]]},
        ]

        blocks, warnings = self.runtime._map_pdf_boxes(boxes, 2, [])

        by_text = {block["text"]: block for block in blocks}
        self.assertEqual("HEADING", by_text["4.1. Payment Elections."]["kind"])
        self.assertEqual("HEADING", by_text["4.3. Deferred Cash."]["kind"])
        self.assertIn(
            "For each director, the company records deferred cash.",
            by_text,
        )
        self.assertIn(
            "DEEPDOC_INLINE_NUMBERED_HEADING_RECONSTRUCTED",
            {warning["code"] for warning in warnings},
        )

    def test_reconstructs_reversed_numbered_heading_fragment(self):
        boxes = [
            {"text": "Document title", "layout_type": "title", "positions": [[1, 0, 90, 0, 8]]},
            {"text": "Duration of Payment Elections.", "layout_type": "title", "positions": [[2, 28, 90, 20, 28]]},
            {"text": "4.2.", "layout_type": "title", "positions": [[2, 10, 22, 20, 28]]},
            {"text": "General Rule. The body remains searchable.", "layout_type": "text", "positions": [[2, 10, 95, 55, 65]]},
        ]

        blocks, warnings = self.runtime._map_pdf_boxes(boxes, 2, [])

        by_text = {block["text"]: block for block in blocks}
        self.assertEqual(
            "HEADING",
            by_text["4.2. Duration of Payment Elections."]["kind"],
        )
        self.assertIn("General Rule. The body remains searchable.", by_text)
        self.assertIn(
            "DEEPDOC_INLINE_NUMBERED_HEADING_RECONSTRUCTED",
            {warning["code"] for warning in warnings},
        )

    def test_does_not_reconstruct_simple_numbered_body_from_text_fragment(self):
        boxes = [
            {"text": "Document title", "layout_type": "title", "positions": [[1, 0, 90, 0, 8]]},
            {"text": "1.", "layout_type": "title", "positions": [[2, 10, 22, 20, 28]]},
            {"text": "Payment must be made within thirty days.", "layout_type": "text", "positions": [[2, 28, 95, 20, 28]]},
        ]

        blocks, warnings = self.runtime._map_pdf_boxes(boxes, 2, [])

        self.assertEqual("1.", blocks[1]["text"])
        self.assertEqual("Payment must be made within thirty days.", blocks[2]["text"])
        self.assertNotIn(
            "DEEPDOC_INLINE_NUMBERED_HEADING_RECONSTRUCTED",
            {warning["code"] for warning in warnings},
        )

    def test_normalizes_reversed_numbered_title_and_splits_heading_prefix(self):
        boxes = [
            {"text": "Document title", "layout_type": "title", "positions": [[1, 0, 90, 0, 8]]},
            {"text": "Duration of Payment Elections.4.2.", "layout_type": "title", "positions": [[2, 10, 90, 20, 28]]},
            {"text": "General Rule. Body text for section 4.2 remains available.", "layout_type": "text", "positions": [[2, 10, 90, 34, 44]]},
            {"text": "4.3. Deferred Cash. For each director, the ﬁrm records deferred cash.", "layout_type": "text", "positions": [[2, 10, 90, 55, 70]]},
        ]

        blocks, warnings = self.runtime._map_pdf_boxes(boxes, 2, [])

        by_text = {block["text"]: block for block in blocks}
        self.assertEqual("HEADING", by_text["4.2. Duration of Payment Elections."]["kind"])
        self.assertEqual("HEADING", by_text["4.3. Deferred Cash."]["kind"])
        self.assertIn(
            "For each director, the ﬁrm records deferred cash.",
            by_text,
        )
        codes = {warning["code"] for warning in warnings}
        self.assertIn("DEEPDOC_NUMBERED_TITLE_ORDER_NORMALIZED", codes)
        self.assertIn("DEEPDOC_NUMBERED_PARAGRAPH_HEADING_SPLIT", codes)

    def test_does_not_promote_numbered_sentence_prefix_from_body_text(self):
        boxes = [
            {"text": "Document title", "layout_type": "title", "positions": [[1, 0, 90, 0, 8]]},
            {"text": "4.3. This agreement is effective today. Remaining contract text.", "layout_type": "text", "positions": [[2, 10, 90, 20, 35]]},
        ]

        blocks, warnings = self.runtime._map_pdf_boxes(boxes, 2, [])

        self.assertEqual(2, len(blocks))
        self.assertEqual("PARAGRAPH", blocks[1]["kind"])
        self.assertEqual(
            "4.3. This agreement is effective today. Remaining contract text.",
            blocks[1]["text"],
        )
        self.assertNotIn(
            "DEEPDOC_NUMBERED_PARAGRAPH_HEADING_SPLIT",
            {warning["code"] for warning in warnings},
        )

    def test_normalizes_glued_numbered_heading_before_text_preservation_hash(self):
        boxes = [
            {"text": "Document title", "layout_type": "title", "positions": [[1, 0, 90, 0, 8]]},
            {"text": "2.2.Medium of Payment. Unless otherwise determined by the Board.", "layout_type": "text", "positions": [[2, 10, 90, 20, 35]]},
        ]

        blocks, warnings = self.runtime._map_pdf_boxes(boxes, 2, [])

        by_text = {block["text"]: block for block in blocks}
        self.assertEqual("HEADING", by_text["2.2. Medium of Payment."]["kind"])
        self.assertIn("Unless otherwise determined by the Board.", by_text)
        codes = {warning["code"] for warning in warnings}
        self.assertIn("DEEPDOC_NUMBERED_PARAGRAPH_HEADING_SPLIT", codes)
        self.assertIn("DEEPDOC_TITLE_GOVERNANCE_TEXT_PRESERVED", codes)

    def test_only_promotes_chapter_level_visual_titles(self):
        boxes = [
            {"text": "Document title", "layout_type": "title", "positions": [[1, 0, 90, 0, 8]]},
            {"text": "Board Agenda", "layout_type": "title", "positions": [[1, 0, 90, 20, 28]]},
            {"text": "Agenda body text is substantial but the label is not a chapter.", "layout_type": "text", "positions": [[1, 0, 90, 35, 45]]},
            {"text": "INTRODUCTION", "layout_type": "title", "positions": [[1, 0, 90, 55, 63]]},
            {"text": "Introduction body text is substantial enough for navigation.", "layout_type": "text", "positions": [[1, 0, 90, 70, 80]]},
            {"text": "结论", "layout_type": "title", "positions": [[1, 0, 90, 90, 98]]},
            {"text": "结论正文提供足够的章节内容用于导航和检索。", "layout_type": "text", "positions": [[1, 0, 90, 105, 115]]},
        ]

        blocks, warnings = self.runtime._map_pdf_boxes(boxes, 1, [])

        by_text = {block["text"]: block for block in blocks}
        self.assertEqual("PARAGRAPH", by_text["Board Agenda"]["kind"])
        self.assertEqual("HEADING", by_text["INTRODUCTION"]["kind"])
        self.assertEqual("HEADING", by_text["结论"]["kind"])
        self.assertIn(
            "DEEPDOC_WEAK_VISUAL_TITLE_DOWNGRADED",
            {warning["code"] for warning in warnings},
        )

    def test_numbered_structure_suppresses_untrusted_uppercase_labels(self):
        boxes = [
            {"text": "Document title", "layout_type": "title", "positions": [[1, 0, 90, 0, 8]]},
        ]
        for index in range(1, 5):
            boxes.extend([
                {"text": f"UNIT {index}: Chapter {index}", "layout_type": "title", "positions": [[index, 0, 90, 20, 28]]},
                {"text": f"Chapter {index} body text with enough content for retrieval navigation.", "layout_type": "text", "positions": [[index, 0, 90, 35, 45]]},
            ])
        boxes.extend([
            {"text": "PERSONAL REFLECTION", "layout_type": "title", "positions": [[5, 0, 90, 20, 28]]},
            {"text": "Reflection activity text remains searchable as ordinary evidence.", "layout_type": "text", "positions": [[5, 0, 90, 35, 45]]},
        ])

        blocks, _warnings = self.runtime._map_pdf_boxes(boxes, 5, [])

        headings = [block["text"] for block in blocks if block["kind"] == "HEADING"]
        self.assertEqual(
            [f"UNIT {index}: Chapter {index}" for index in range(1, 5)],
            headings,
        )
        reflection = next(
            block for block in blocks if block["text"] == "PERSONAL REFLECTION"
        )
        self.assertEqual("PARAGRAPH", reflection["kind"])

    def test_docx_interleaves_headings_paragraphs_and_tables_in_body_order(self):
        document = Document()
        document.add_heading("Section A", 1)
        table_a = document.add_table(rows=2, cols=1)
        table_a.cell(0, 0).text = "A first row"
        table_a.cell(1, 0).text = "A second row"
        document.add_paragraph("Section A body")
        document.add_heading("Section B", 1)
        table_b = document.add_table(rows=2, cols=1)
        table_b.cell(0, 0).text = "B first row"
        table_b.cell(1, 0).text = "B second row"
        document.add_paragraph("Section B body")
        parser = MagicMock(doc=document)
        parser.return_value = (
            [(paragraph.text, paragraph.style.name) for paragraph in document.paragraphs],
            [
                [cell.text for row in table.rows for cell in row.cells]
                for table in document.tables
            ],
        )
        self.runtime.parsers["DOCX"] = parser

        blocks, warnings = self.runtime._parse_docx(Path("unused.docx"))

        self.assertEqual(
            ["Section A", "A first row", "A second row", "Section A body",
             "Section B", "B first row", "B second row", "Section B body"],
            [block["text"] for block in blocks],
        )
        self.assertEqual(
            [0, 1, 1, 2, 3, 4, 4, 5],
            [block["bodyElementIndex"] for block in blocks],
        )
        self.assertEqual(
            ["HEADING", "TABLE_CELL", "TABLE_CELL", "PARAGRAPH",
             "HEADING", "TABLE_CELL", "TABLE_CELL", "PARAGRAPH"],
            [block["kind"] for block in blocks],
        )
        self.assertNotIn(
            "DEEPDOC_DOCX_POSITION_LOGICAL",
            {warning["code"] for warning in warnings},
        )

    def test_docx_recovers_single_row_table_without_shifting_or_duplicating_later_table(self):
        document = Document()
        document.add_heading("Section A", 1)
        document.add_table(rows=1, cols=1).cell(0, 0).text = "ZephyrQuota is 42 units."
        document.add_heading("Section B", 1)
        later_table = document.add_table(rows=2, cols=1)
        later_table.cell(0, 0).text = "Quota"
        later_table.cell(1, 0).text = "84 units"
        document.add_paragraph("Closing notes")
        parser = MagicMock(doc=document)
        parser.return_value = (
            [(paragraph.text, paragraph.style.name) for paragraph in document.paragraphs],
            [[], ["Quota: 84 units"]],
        )
        self.runtime.parsers["DOCX"] = parser

        blocks, warnings = self.runtime._parse_docx(Path("unused.docx"))

        self.assertEqual(
            ["Section A", "ZephyrQuota is 42 units.", "Section B", "Quota: 84 units", "Closing notes"],
            [block["text"] for block in blocks],
        )
        self.assertEqual([0, 1, 2, 3, 4], [block["bodyElementIndex"] for block in blocks])
        self.assertEqual(
            ["HEADING", "TABLE_CELL", "HEADING", "TABLE_CELL", "PARAGRAPH"],
            [block["kind"] for block in blocks],
        )
        self.assertNotIn(
            "DEEPDOC_DOCX_POSITION_LOGICAL",
            {warning["code"] for warning in warnings},
        )

    def test_docx_empty_table_fallback_deduplicates_merged_cells_by_identity_not_text(self):
        document = Document()
        document.add_heading("Section A", 1)
        table = document.add_table(rows=1, cols=3)
        table.cell(0, 0).merge(table.cell(0, 1)).text = "Same quota"
        table.cell(0, 2).text = "Same quota"
        parser = MagicMock(doc=document)
        parser.return_value = ([("Section A", "Heading 1")], [[]])
        self.runtime.parsers["DOCX"] = parser

        blocks, _warnings = self.runtime._parse_docx(Path("unused.docx"))

        self.assertEqual(["Section A", "Same quota", "Same quota"], [block["text"] for block in blocks])
        self.assertEqual([0, 1, 1], [block["bodyElementIndex"] for block in blocks])

    def test_docx_empty_table_fallback_does_not_invent_text_without_source_or_content(self):
        empty_document = Document()
        empty_document.add_heading("Section A", 1)
        empty_document.add_table(rows=1, cols=1)
        for document in (None, empty_document):
            with self.subTest(has_document=document is not None):
                parser = MagicMock(doc=document)
                parser.return_value = ([("Section A", "Heading 1")], [[]])
                self.runtime.parsers["DOCX"] = parser

                blocks, _warnings = self.runtime._parse_docx(Path("unused.docx"))

                self.assertEqual(["Section A"], [block["text"] for block in blocks])

    def test_docx_empty_table_fallback_rejects_unaligned_table_ordinals(self):
        document = Document()
        document.add_heading("Section A", 1)
        document.add_table(rows=1, cols=1).cell(0, 0).text = "First table"
        document.add_table(rows=1, cols=1).cell(0, 0).text = "Second table"
        parser = MagicMock(doc=document)
        # A parser that omits empty placeholders cannot safely identify this table.
        parser.return_value = ([("Section A", "Heading 1")], [[]])
        self.runtime.parsers["DOCX"] = parser

        blocks, _warnings = self.runtime._parse_docx(Path("unused.docx"))

        self.assertEqual(["Section A"], [block["text"] for block in blocks])

    def test_docx_without_document_positions_retains_logical_order_and_warning(self):
        blocks, warnings = self.runtime._parse_docx(Path("unused.docx"))

        self.assertEqual(
            ["Document title", "Section", "Body", "Name: A;Value: 1"],
            [block["text"] for block in blocks],
        )
        self.assertEqual([0, 1, 2, 3], [block["bodyElementIndex"] for block in blocks])
        self.assertIn(
            "DEEPDOC_DOCX_POSITION_LOGICAL",
            {warning["code"] for warning in warnings},
        )

    def test_dispatches_docx_txt_and_markdown_with_native_positions(self):
        docx = self.post(docx_bytes(), "DOCX")
        self.assertEqual(200, docx.status_code)
        docx_body = docx.get_json()
        self.assertEqual("DOCX", docx_body["sourceFormat"])
        self.assertEqual("TITLE", docx_body["blocks"][0]["kind"])
        self.assertEqual("HEADING", docx_body["blocks"][1]["kind"])
        self.assertEqual(2, docx_body["blocks"][1]["headingLevel"])
        self.assertEqual("TABLE_CELL", docx_body["blocks"][-1]["kind"])

        text = self.post(b"first\nsecond\n\nthird", "TXT")
        self.assertEqual(200, text.status_code)
        text_blocks = text.get_json()["blocks"]
        self.assertEqual((1, 2), (text_blocks[0]["startLine"], text_blocks[0]["endLine"]))
        self.assertEqual((4, 4), (text_blocks[1]["startLine"], text_blocks[1]["endLine"]))

        markdown = self.post(b"# Heading\nBody", "MARKDOWN")
        self.assertEqual(200, markdown.status_code)
        markdown_blocks = markdown.get_json()["blocks"]
        self.assertEqual("HEADING", markdown_blocks[0]["kind"])
        self.assertEqual("Heading", markdown_blocks[0]["text"])
        self.assertEqual(1, markdown_blocks[0]["startLine"])

    def test_digest_mismatch_is_permanent_client_error(self):
        with self.app.test_client() as client:
            response = client.post(
                "/v1/parse",
                data={"file": (io.BytesIO(self.payload), "document.pdf")},
                headers={
                    "X-Source-SHA256": "0" * 64,
                    "X-Source-Format": "PDF",
                },
                content_type="multipart/form-data",
            )
        self.assertEqual(400, response.status_code)
        self.assertEqual(
            "SOURCE_OBJECT_INTEGRITY_MISMATCH",
            response.get_json()["error"]["code"],
        )
        self.assertFalse(response.get_json()["error"]["retryable"])


if __name__ == "__main__":
    unittest.main()
