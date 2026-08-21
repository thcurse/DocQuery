"""Build the binary documents and checksummed manifest for n3-eval-v1.

The corpus is wholly fictional. Run with the Codex bundled Python runtime.
"""

from __future__ import annotations

import hashlib
import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer


ROOT = Path(__file__).resolve().parents[2]
EVAL_ROOT = ROOT / "evaluation" / "n3-eval-v1"
CORPUS = EVAL_ROOT / "corpus"


def set_font(run, ascii_font="Calibri", cjk_font="Microsoft YaHei", size=None,
             bold=None, color=None):
    run.font.name = ascii_font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), ascii_font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), ascii_font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), cjk_font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def configure_docx(document: Document, title: str):
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, before, after, color in (
        ("Heading 1", 16, 18, 10, (46, 116, 181)),
        ("Heading 2", 13, 14, 7, (46, 116, 181)),
        ("Heading 3", 12, 10, 5, (31, 77, 120)),
    ):
        style = document.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(*color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    title_paragraph = document.add_paragraph()
    title_paragraph.paragraph_format.space_after = Pt(8)
    run = title_paragraph.add_run(title)
    set_font(run, size=24, bold=True, color=(11, 37, 69))

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(18)
    run = subtitle.add_run("虚构资料 - DocQuery N3 评测专用")
    set_font(run, size=10, color=(90, 100, 112))

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("DocQuery n3-eval-v1")
    set_font(run, size=9, color=(110, 118, 126))


def add_heading(document: Document, text: str, level: int):
    document.add_heading(text, level=level)


def add_body(document: Document, text: str):
    paragraph = document.add_paragraph(text)
    paragraph.paragraph_format.keep_together = False


def build_expense_policy():
    doc = Document()
    configure_docx(doc, "澄岳差旅与费用报销制度")
    add_heading(doc, "交通与住宿", 1)
    add_body(doc, "高铁行程不超过五小时的，原则上购买二等座。单晚住宿含税上限为680元；北京、上海、深圳三地在展会期间可以上浮至860元，但报销时必须附会议通知。")
    add_heading(doc, "餐费与招待", 1)
    add_body(doc, "出差餐费按每人每天160元包干。已经由客户或会议主办方提供正餐的，当日包干额扣减50元。业务招待费不与餐费包干重复报销。")
    add_heading(doc, "预审批", 1)
    add_body(doc, "预计总费用超过5000元的差旅必须在出发前取得直属经理批准。临时抢修任务可以先出发，但应在任务开始后24小时内补交预审批。")
    add_heading(doc, "票据与提交期限", 1)
    add_body(doc, "报销申请应在返程后10个工作日内提交。电子发票必须上传原始PDF文件，截图不能替代原始发票。")
    add_heading(doc, "例外处理", 1)
    add_body(doc, "因航班取消导致住宿超过标准时，需要同时提交航空公司取消证明和酒店结算单。只有其中一项材料时，超标部分暂不报销。")
    doc.save(CORPUS / "travel-expense-policy.docx")


def build_security_standard():
    doc = Document()
    configure_docx(doc, "凌光生产系统访问安全标准")
    add_heading(doc, "账号与认证", 1)
    add_heading(doc, "登录失败", 2)
    add_body(doc, "同一账号连续五次密码失败后锁定30分钟。安全管理员可以提前解锁，但必须填写事件编号；普通平台主管不能提前解锁。")
    add_heading(doc, "多因素认证", 2)
    add_body(doc, "生产控制台和密钥管理页面必须使用多因素认证。只读监控大屏使用设备证书，不允许共享个人验证码。")
    add_heading(doc, "紧急访问", 1)
    add_body(doc, "紧急账号启用需要值班经理和安全负责人双人批准，有效期最长两小时。使用结束后15分钟内必须自动禁用，并在次日12:00前完成访问复盘。")
    add_heading(doc, "凭证轮换", 1)
    add_body(doc, "生产应用密钥每90天轮换一次。新旧密钥最多并存24小时；确认新密钥成功调用后，应立即撤销旧密钥。")
    add_heading(doc, "审计要求", 1)
    add_body(doc, "审计记录可以包含应用、资源、结果和不透明actorRef，不得把actorRef当作可信用户身份。Authorization和完整Secret不得写入日志或审计。")
    doc.save(CORPUS / "security-access-standard.docx")


def build_supplier_returns():
    doc = Document()
    configure_docx(doc, "星岚供应商退货作业指导书")
    add_heading(doc, "退货受理", 1)
    add_body(doc, "发现来料缺陷后两个工作日内创建RMA申请。申请必须包含采购订单号、物料批次、缺陷照片和抽检数量。")
    add_heading(doc, "包装与标签", 1)
    add_body(doc, "退货外箱使用橙色RMA标签。每箱只能包含一个物料批次，标签上的RMA编号必须与装箱单一致。静电敏感器件还必须使用银色防静电袋。")
    add_heading(doc, "承运安排", 1)
    add_body(doc, "供应商承担质量缺陷退货运费。仓库在收到供应商取件码后才能交付承运人；不得使用采购员口头提供的临时地址。")
    add_heading(doc, "账务关闭", 1)
    add_body(doc, "供应商签收后，采购员在三个工作日内上传签收证明。财务收到签收证明和红字发票后关闭RMA；缺少任一材料时保持待结算状态。")
    add_heading(doc, "数量差异", 1)
    add_body(doc, "供应商签收数量与装箱单不一致时，仓库不得直接修改原装箱单，应创建差异记录并保留出库称重照片。")
    doc.save(CORPUS / "supplier-return-process.docx")


def pdf_styles():
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    base = ParagraphStyle(
        "CJKBody", fontName="STSong-Light", fontSize=10.5, leading=17,
        textColor=HexColor("#202B36"), spaceAfter=9,
    )
    title = ParagraphStyle(
        "CJKTitle", parent=base, fontSize=22, leading=28,
        alignment=TA_CENTER, textColor=HexColor("#0B2545"), spaceAfter=20,
    )
    h1 = ParagraphStyle(
        "CJKH1", parent=base, fontSize=16, leading=22,
        textColor=HexColor("#1F4D78"), spaceBefore=15, spaceAfter=9,
    )
    h2 = ParagraphStyle(
        "CJKH2", parent=base, fontSize=13, leading=19,
        textColor=HexColor("#2E74B5"), spaceBefore=12, spaceAfter=7,
    )
    return base, title, h1, h2


def pdf_doc(path: Path, story):
    document = SimpleDocTemplate(
        str(path), pagesize=letter,
        rightMargin=0.85 * inch, leftMargin=0.85 * inch,
        topMargin=0.8 * inch, bottomMargin=0.8 * inch,
        title=path.stem, author="DocQuery fictional evaluation corpus",
    )
    document.build(story)


def build_mx7_manual():
    body, title, _, _ = pdf_styles()
    pages = [
        [
            "MX7循环冷却机是星岚实验线的封闭式冷却设备。本手册中的设备编号、阀门名称和阈值均为评测使用的虚构资料。",
            "设备每天开机前检查储液罐刻度。液位低于MIN线时只能补充型号CL-8冷却液，不得混入去离子水。环境温度低于5摄氏度时，先让预热器运行12分钟再启动主泵。",
            "正常运行时入口压力为1.6至2.1巴，出口温度为18至24摄氏度。观察窗出现连续气泡时，先检查回液软管是否折弯，再检查过滤杯密封圈。",
        ],
        [
            "故障E17表示冷却回路压力建立失败。复位顺序是隔离加工线、等待90秒、完全打开旁通阀BV-4，然后按住RST键6秒。每小时不得复位超过两次；第三次出现E17时必须停机检查泵入口。",
            "复位后先空载运行三分钟。若入口压力仍低于1.4巴，检查过滤杯和泵前软管，不得通过提高目标转速掩盖压力不足。",
            "维护人员把E17俗称为“水压起不来”或“冷却泵憋住了”，但工单中必须保留正式错误码E17。",
        ],
        [
            "告警F42表示主泵振动速度达到7.2毫米每秒或更高。处理时检查C-19联轴器的紧固标记和弹性体，不得先更换压力传感器。振动恢复到5.0毫米每秒以下并持续十分钟后才能解除告警。",
            "如果F42与E17同时出现，先处理F42并保持设备停机，完成联轴器检查后再执行E17复位流程。两个故障的先后顺序不能互换。",
            "季度维护需要校验振动探头零点，允许偏差为正负0.15毫米每秒。超差探头送计量室，不在现场修改校准系数。",
        ],
        [
            "过滤芯型号为FX-22，正常每800运行小时更换。压差超过0.35巴时即使未满800小时也要更换。更换前关闭V1和V2并释放杯体压力，安装后用手拧紧四分之一圈。",
            "废过滤芯放入蓝色密闭桶，标签注明设备编号、运行小时和更换日期。冷却液渗漏超过200毫升时，使用吸液垫处理并创建ENV-SPILL记录。",
            "过滤芯更换完成后记录入口压力、出口温度和滤杯是否漏液。缺少任一项读数时维护记录不完整。",
        ],
        [
            "长期停机超过30天时排空冷却液，用氮气以0.3巴压力吹扫回路两分钟。重新启用时先进行密封测试，不得直接连接加工线。",
            "运输前锁定压缩机支架并拆下储液罐液位探头。运输倾斜角不得超过25度，到场后静置四小时才能通电。",
            "设备外壳只能使用中性清洁剂。含氯溶剂会损坏观察窗和软管标识，禁止用于日常清洁。",
        ],
        [
            "维护记录至少保留三年。记录包括设备序列号、错误码、执行步骤、替换零件和复测结果，不记录操作人员的个人手机号。",
            "远程支持需要requestId和设备序列号后四位。禁止上传包含客户生产配方的控制器完整导出文件。",
            "本设备没有自动补液功能，也不支持手机蓝牙配置。任何宣称可以通过蓝牙解除E17的说明都不属于本手册。",
        ],
    ]
    story = [Paragraph("MX7循环冷却机维护手册", title)]
    for index, paragraphs in enumerate(pages):
        if index:
            story.append(PageBreak())
        for paragraph in paragraphs:
            story.append(Paragraph(paragraph, body))
            story.append(Spacer(1, 4))
    pdf_doc(CORPUS / "mx7-maintenance-manual.pdf", story)


def build_a1_guide():
    body, title, h1, _ = pdf_styles()
    sections = [
        ("设备登记", [
            "A1风机首次投入运行前必须登记叶轮序列号和控制器固件版本。登记完成后生成A1-REG编号，缺少该编号不能申请保修。",
            "设备转移到其他站点时保留原序列号，但要在48小时内更新站点代码。",
        ]),
        ("例行检查", [
            "每500运行小时检查皮带张力。指针应落在绿色区域3至5格；低于3格时调整张紧轮，高于5格时释放张力。",
            "进风滤网每30天清洁一次。粉尘红色预警站点缩短为每14天一次。",
        ]),
        ("过热停机", [
            "告警A1-HOT在电机绕组温度达到118摄氏度时触发。停机后至少等待20分钟，并确认温度低于70摄氏度，才能执行复位。",
            "若同一自然日发生两次A1-HOT，第二次不得复位，应检查风道和皮带打滑。",
        ]),
        ("保修范围", [
            "标准保修期为验收之日起18个月。皮带和滤网属于耗材，不在标准保修范围；因控制器制造缺陷导致的更换包含人工费。",
            "未登记A1-REG编号、擅自修改过热阈值或使用非原厂叶轮会终止相关故障的保修。",
        ]),
        ("远程支持", [
            "远程支持前导出最近15分钟趋势数据，只包含转速、温度和振动。客户配方和操作员姓名不得进入导出文件。",
        ]),
    ]
    story = [Paragraph("A1工业风机服务指南", title)]
    for heading, paragraphs in sections:
        story.append(Paragraph(heading, h1))
        for paragraph in paragraphs:
            story.append(Paragraph(paragraph, body))
    pdf_doc(CORPUS / "a1-service-guide.pdf", story)


def build_incident_guide():
    body, title, h1, h2 = pdf_styles()
    story = [Paragraph("澄岳联合事件响应指南", title)]
    content = [
        ("事件分级", [
            ("一级事件", "核心交易完全不可用且影响两个以上区域，或确认发生敏感数据外泄，定义为一级事件。值班人员必须在10分钟内通知事件指挥官。"),
            ("二级事件", "单一区域核心功能降级超过20分钟但存在人工替代方案，定义为二级事件。"),
        ]),
        ("指挥与沟通", [
            ("指挥权", "一级事件由当周事件指挥官统一决策。平台主管负责技术建议，但不得绕过指挥官直接发布客户公告。"),
            ("对外更新", "一级事件首次客户更新应在确认后30分钟内发出，此后每45分钟更新一次，即使没有新的恢复进展也要说明当前状态。"),
        ]),
        ("证据保护", [
            ("日志快照", "涉及疑似入侵时，先创建只读日志快照并记录SHA-256，再执行主机隔离。不得为了节省时间先清空临时目录。"),
            ("访问控制", "证据对象只允许事件指挥官、安全负责人和法务指定人员读取。普通平台主管没有默认读取权限。"),
        ]),
        ("恢复与关闭", [
            ("恢复门禁", "恢复生产流量前必须同时满足错误率连续15分钟低于1%、关键队列无持续增长、值班经理批准三个条件。"),
            ("复盘", "一级事件在恢复后两个工作日内完成初稿，五个工作日内完成正式复盘。正式复盘必须列出时间线、根因、客户影响和责任明确的改进项。"),
        ]),
    ]
    for heading, subsections in content:
        story.append(Paragraph(heading, h1))
        for subheading, paragraph in subsections:
            story.append(Paragraph(subheading, h2))
            story.append(Paragraph(paragraph, body))
    pdf_doc(CORPUS / "incident-response-guide.pdf", story)


def write_manifest():
    documents = [
        ("mx7-manual", "mx7-maintenance-manual.pdf", "PDF", "MX7循环冷却机维护手册"),
        ("a1-service", "a1-service-guide.pdf", "PDF", "A1工业风机服务指南"),
        ("incident-guide", "incident-response-guide.pdf", "PDF", "澄岳联合事件响应指南"),
        ("expense-policy", "travel-expense-policy.docx", "DOCX", "澄岳差旅与费用报销制度"),
        ("security-standard", "security-access-standard.docx", "DOCX", "凌光生产系统访问安全标准"),
        ("supplier-returns", "supplier-return-process.docx", "DOCX", "星岚供应商退货作业指导书"),
        ("orion-api", "orion-api-operations.md", "MARKDOWN", "Orion API 运维手册"),
        ("refund-rules", "customer-refund-rules.md", "MARKDOWN", "星澜客户退款规则"),
        ("retention-handbook", "data-retention-handbook.md", "MARKDOWN", "澄岳数据保留手册"),
        ("warehouse-shift", "warehouse-shift-procedures.txt", "TXT", "澄岳仓库交接班操作说明"),
        ("battery-guide", "battery-inspection-log-guide.txt", "TXT", "凌光电池巡检记录指南"),
        ("field-pricing", "field-service-pricing.txt", "TXT", "星澜现场服务计价说明"),
    ]
    manifest_docs = []
    checksum_lines = []
    for key, filename, source_format, display_name in documents:
        data = (CORPUS / filename).read_bytes()
        digest = hashlib.sha256(data).hexdigest()
        manifest_docs.append({
            "documentKey": key,
            "file": f"corpus/{filename}",
            "format": source_format,
            "displayName": display_name,
            "sha256": digest,
            "bytes": len(data),
        })
        checksum_lines.append(f"{digest}  corpus/{filename}")
    manifest = {
        "datasetVersion": "n3-eval-v1",
        "fictional": True,
        "documentCount": len(manifest_docs),
        "caseCount": 40,
        "developmentCaseCount": 24,
        "acceptanceCaseCount": 16,
        "documents": manifest_docs,
    }
    (EVAL_ROOT / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    (EVAL_ROOT / "checksums.sha256").write_text(
        "\n".join(checksum_lines) + "\n", encoding="ascii"
    )


def main():
    CORPUS.mkdir(parents=True, exist_ok=True)
    build_mx7_manual()
    build_a1_guide()
    build_incident_guide()
    build_expense_policy()
    build_security_standard()
    build_supplier_returns()
    write_manifest()


if __name__ == "__main__":
    main()
